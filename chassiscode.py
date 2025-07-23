import pandas as pd
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.chrome.service import Service
from webdriver_manager.chrome import ChromeDriverManager
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.support.ui import WebDriverWait
import time
from openpyxl import load_workbook
import random
FILEPATH = "split_3"

df = pd.read_excel(FILEPATH+".xlsx", header=None)
wb = load_workbook(FILEPATH+".xlsx")
ws = wb.active

#setup
options = webdriver.ChromeOptions()
#udaan bololtoi
# options.add_argument('start-maximized')
#hurdan bololtoi
options.add_argument('--headless')
options.add_argument('--disable-gpu')
service = Service(ChromeDriverManager().install())
driver = webdriver.Chrome(service=service, options=options)

#logic
save_interval = 50

for index, row in df.iterrows():
    VIN = row.iloc[0]
    print(VIN)
    if not ws.cell(row=index + 1, column=2).value:
        start_time = time.time()
        driver.get("https://www.toyodiy.com/parts/q?vin="+str(VIN))
        time.sleep(0.5)
        try:
            production_date = driver.find_element(By.XPATH, '//a[@title="production date"]').text.strip()
            row_xpath = '//table[@class="res"]/tbody/tr[2]' 
            model_code = driver.find_element(By.XPATH, row_xpath + '/td[2]').text.strip()
            from_date = driver.find_element(By.XPATH, row_xpath + '/td[3]').text.strip()
            to_date = driver.find_element(By.XPATH, row_xpath + '/td[4]').text.strip()

            print("Model Code:", model_code)
            print("From Date:", from_date)
            print("To Date:", to_date)

            ws.cell(row=index + 1, column=2, value=model_code)  
            ws.cell(row=index + 1, column=3, value=from_date)   
            ws.cell(row=index + 1, column=4, value=to_date) 
            ws.cell(row=index + 1, column=5, value=production_date)    
        except Exception as e:
            print("Aldaa")
            ws.cell(row = index + 1, column = 2, value = "Not found")
            
        if index % save_interval == 0:
            wb.save(FILEPATH+".xlsx")
            
        end_time = time.time()
        hugatsaa = end_time - start_time
        print(f"hugatsaa {hugatsaa:.3f}")
        print(f"Row: {index}")
    else:
        print(f"Row {index} processed")
        continue
    
wb.save(FILEPATH+".xlsx")
driver.quit()
print("Amjilttai")