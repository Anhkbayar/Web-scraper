import pandas as pd
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.chrome.service import Service
from webdriver_manager.chrome import ChromeDriverManager
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.support.ui import WebDriverWait
import time
from openpyxl import load_workbook
import random
from openpyxl.styles import PatternFill

FILEPATH = "split_1"

EMAIL = "ankhbayar@garage.mn"
PASSWORD = "Ankhaa#123"
RED_FILL = PatternFill(start_color='FF0000', end_color='FF0000', fill_type='solid')
GREEN_FILL = PatternFill(start_color="11AD30", end_color='C6EFCE', fill_type='solid')

df = pd.read_excel(FILEPATH+".xlsx")
wb = load_workbook(FILEPATH+".xlsx")
ws = wb["Sheet5"]

#setup
options = webdriver.ChromeOptions()
options.add_argument('start-maximized')
service = Service(ChromeDriverManager().install())
driver = webdriver.Chrome(service=service, options=options)

#login
driver.get("https://gp.garage.mn/parts")
time.sleep(1)

loginbttn = driver.find_element(By.ID, "btnlogin")
loginbttn.click()

driver.find_element(By.ID, "email").send_keys(EMAIL)
driver.find_element(By.ID, "password").send_keys(PASSWORD, Keys.RETURN)

for index, row in df.iterrows():
    driver.get("https://gp.garage.mn/parts")
    time.sleep(1)
    
    driver.find_element(By.ID, "articleid").send_keys(str(article_id), Keys.RETURN)
    time.sleep(1)
    
    link = WebDriverWait(driver, 10).until(
        EC.element_to_be_clickable((By.CSS_SELECTOR, "a.text-body.fw-bold"))
    )