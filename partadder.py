import pandas as pd
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.chrome.service import Service
from webdriver_manager.chrome import ChromeDriverManager
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.support.ui import WebDriverWait
import time
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

EMAIL = "ankhbayar@garage.mn"
PASSWORD = "Ankhaa#123"
RED_FILL = PatternFill(start_color='FF0000', end_color='FF0000', fill_type='solid')
BRAND = "Hyundai"


#unshih
df = pd.read_excel("TestArticle.xlsx")
wb = load_workbook("TestArticle.xlsx")
    
options = webdriver.ChromeOptions()
options.add_argument('--headless')
options.add_argument('--disable-gpu')
options.add_argument('--no-sandbox')
options.add_argument('--disable-dev-shm-usage')
    
service = Service(ChromeDriverManager().install())
driver = webdriver.Chrome(service=service, options=options)
# driver.set_page_load_timeout(page_load_timeout)

driver.get("https://gp.garage.mn/parts")
time.sleep(1)

loginbttn = driver.find_element(By.ID, "btnlogin")
loginbttn.click()

driver.find_element(By.ID, "email").send_keys(EMAIL)
driver.find_element(By.ID, "password").send_keys(PASSWORD, Keys.RETURN)

for index, row in df.iterrows():
    Angilal = row.iloc[6]
    Ded_angilal = row.iloc[7]
    Generic_article = row.iloc[8]
    
    driver.get("https://gp.garage.mn/parts")
    time.sleep(1)
    
    
    
    

