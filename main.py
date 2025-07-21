import pandas as pd
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.chrome.service import Service
from webdriver_manager.chrome import ChromeDriverManager
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.support.ui import WebDriverWait
import time
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

EMAIL = "ankhbayar@garage.mn"
PASSWORD = "Ankhaa#123"
RED_FILL = PatternFill(start_color='FF0000', end_color='FF0000', fill_type='solid')


#unshih
df = pd.read_excel("TestArticle.xlsx")
wb = load_workbook("TestArticle.xlsx")

id_sheet = wb.worksheets[0]
fitment_sheet = wb["Fitment Data"] if "Fitment Data" in wb.sheetnames else wb.create_sheet("Fitment Data")

#setup
options = webdriver.ChromeOptions()
options.add_argument('start-maximized')
service = Service(ChromeDriverManager().install())
driver = webdriver.Chrome(service=service, options=options)

#login
driver.get("https://gp.garage.mn/parts")
time.sleep(1)

loginbttn = driver.find_element(By.ID, "btnlogin")
loginbttn.click()

driver.find_element(By.ID, "email").send_keys(EMAIL)
driver.find_element(By.ID, "password").send_keys(PASSWORD, Keys.RETURN)

#process
for index, row in df.iterrows():
    article_id = row['ArticleID']
    
    driver.get("https://gp.garage.mn/parts")
    time.sleep(1)
    
    driver.find_element(By.ID, "articleid").send_keys(str(article_id), Keys.RETURN)
    time.sleep(1)
    
    link = WebDriverWait(driver, 10).until(
        EC.element_to_be_clickable((By.CSS_SELECTOR, "a.text-body.fw-bold"))
    )
    link.click()
    
    #OEM huulah
    try:
        h5_element = WebDriverWait(driver, 5).until(
            EC.presence_of_element_located((By.CSS_SELECTOR, "h5.text-truncate.font-size-14"))
        )

        part_number = h5_element.text.strip()
        print("Huulsan oem:", part_number)
        
        driver.get("https://www.toyodiy.com/parts/xref?s="+part_number+"&mU=on&mE=on&mJ=on&mG=on")
        fitment_table = WebDriverWait(driver, 4).until(
        EC.element_to_be_clickable((By.XPATH, "//a[@title='click to reveal details']"))
        )
        fitment_table.click()
        
        WebDriverWait(driver, 10).until(
        lambda d: "Loading" not in d.find_element(By.CSS_SELECTOR, "tbody#res3311D0").text
        )
        
        table = WebDriverWait(driver, 10).until(
            EC.presence_of_element_located((By.CSS_SELECTOR, "tbody#res3311D0"))
        )
        rows = table.find_elements(By.TAG_NAME, "tr")
            
        #data nemeh
        for row in rows:
            cols = row.find_elements(By.TAG_NAME, "td")
            row_data = [str(article_id)]
            row_data.extend(col.text.strip() for col in cols[1:])
            fitment_sheet.append(row_data)
    except Exception as e:
        id_sheet.cell(row = index+2, column = 1).fill = RED_FILL
        print("Aldaatai articleId", str(row))
        
wb.save("TestArticle.xlsx")
driver.quit()
print("Amjilttai")