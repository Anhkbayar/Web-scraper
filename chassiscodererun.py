import pandas as pd
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.chrome.service import Service
from webdriver_manager.chrome import ChromeDriverManager
import time
from openpyxl import load_workbook
import os
from selenium.common.exceptions import TimeoutException, WebDriverException

# Setup
save_interval = 50
start_num = 1
end_num = 1
max_retries = 3
timeout_delay = 1
page_load_timeout = 120

def setup_driver():
    options = webdriver.ChromeOptions()
    options.add_argument('--headless')
    options.add_argument('--disable-gpu')
    options.add_argument('--no-sandbox')
    options.add_argument('--disable-dev-shm-usage')
    
    service = Service(ChromeDriverManager().install())
    driver = webdriver.Chrome(service=service, options=options)
    driver.set_page_load_timeout(page_load_timeout)
    return driver

def process_vin(driver, vin):
    try:
        driver.get(f"https://www.toyodiy.com/parts/q?vin={str(vin)}")
        time.sleep(0.3)
        
        production_date = driver.find_element(By.XPATH, '//a[@title="production date"]').text.strip()
        row_xpath = '//table[@class="res"]/tbody/tr[2]'
        model_code = driver.find_element(By.XPATH, row_xpath + '/td[2]').text.strip()
        from_date = driver.find_element(By.XPATH, row_xpath + '/td[3]').text.strip()
        to_date = driver.find_element(By.XPATH, row_xpath + '/td[4]').text.strip()
        
        return model_code, from_date, to_date, production_date
    except Exception as e:
        print(f"VIN {vin}: {str(e)}")
        print("ERONHII OLDOOGUIDE XD")
        raise

def main():
    driver = setup_driver()
    
    try:
        file_num = start_num
        while file_num <= end_num:
            FILEPATH = f"split_{file_num}.xlsx"
            
            if not os.path.exists(FILEPATH):
                print(f"{FILEPATH} OLDDDDGUIGUGIUAGIUAGUIA")
                file_num += 1
                continue
            
            print(f"\n{FILEPATH} EHELLEEEE YEYYEE")
            
            df = pd.read_excel(FILEPATH, header=None)
            wb = load_workbook(FILEPATH)
            ws = wb.active

            for index, row in df.iterrows():
                VIN = row.iloc[0]
                if ws.cell(row=index + 1, column=5).value or ws.cell(row=index + 1, column=2).value=='Not found':
                    print(f"Row {index} processed")
                    continue
                
                retry_count = 0
                while retry_count < max_retries:
                    try:
                        start_time = time.time()
                        
                        model_code, from_date, to_date, production_date = process_vin(driver, VIN)
                        
                        ws.cell(row=index + 1, column=2, value=model_code)
                        ws.cell(row=index + 1, column=3, value=from_date)
                        ws.cell(row=index + 1, column=4, value=to_date)
                        ws.cell(row=index + 1, column=5, value=production_date)
                        
                        processing_time = time.time() - start_time
                        print(f"{index+1}")
                        print(f"VIN {VIN} in {processing_time:.2f}s")
                        
                        # hadgalah
                        if index % save_interval == 0:
                            wb.save(FILEPATH)
                        
                        break
                        
                    except TimeoutException:
                        retry_count += 1
                        print(f"Timeout on {VIN}")
                        if retry_count >= max_retries:
                            ws.cell(row=index + 1, column=2, value="Timeout error")
                        else:
                            time.sleep(timeout_delay)
                            driver.quit()
                            driver = setup_driver()
                            
                    except WebDriverException as e:
                        retry_count += 1
                        print(f"WebDriver error: {str(e)} (attempt {retry_count}/{max_retries})")
                        if retry_count >= max_retries:
                            ws.cell(row=index + 1, column=2, value="Driver error")
                        else:
                            time.sleep(timeout_delay)
                            driver.quit()
                            driver = setup_driver()
                            
                    except Exception as e:
                        ws.cell(row=index + 1, column=2, value=f"Error: {str(e)}")
                        break
            
            # Neg file duusaad hadgalah
            wb.save(FILEPATH)
            print(f"{FILEPATH} DUUUSLAAA YEYEYEYAYAYAYE :)))))")
            file_num += 1
            
    finally:
        driver.quit()
        print("DUUUSLAA UEUEUEUEEUEU")

if __name__ == "__main__":
    main()