import pandas as pd
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.chrome.service import Service
from webdriver_manager.chrome import ChromeDriverManager
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.support.ui import WebDriverWait
import time
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

EMAIL = "ankhbayar@garage.mn"
PASSWORD = "Ankhaa#123"
RED_FILL = PatternFill(start_color='FF0000', end_color='FF0000', fill_type='solid')
BRAND = "Hyundai"
FILEPATH = f"{BRAND}.xlsx"


#unshih
df = pd.read_excel(FILEPATH, sheet_name=1, header=None)
wb = load_workbook(FILEPATH)
ws = wb.worksheets[1]
    
options = webdriver.ChromeOptions()
options.add_argument('start_maximized')
# options.add_argument('--headless')
# options.add_argument('--disable-gpu')
# options.add_argument('--no-sandbox')
# options.add_argument('--disable-dev-shm-usage')
    
service = Service(ChromeDriverManager().install())
driver = webdriver.Chrome(service=service, options=options)
# driver.set_page_load_timeout(page_load_timeout)

driver.get("https://gp.garage.mn/parts")
time.sleep(1)

loginbttn = driver.find_element(By.ID, "btnlogin")
loginbttn.click()

driver.find_element(By.ID, "email").send_keys(EMAIL)
driver.find_element(By.ID, "password").send_keys(PASSWORD, Keys.RETURN)

for index, row in df.iterrows():
    if ws.cell(row = index+1, column=11).value:
        article_id = row.iloc[7]
        Generic_ID = int(row.iloc[10])
        
        print(article_id)
        print(Generic_ID)
        
        driver.get("https://gp.garage.mn/parts")
        
        driver.find_element(By.ID, "articleno").send_keys(str(article_id), Keys.RETURN)
        time.sleep(1)
        
        first_p = driver.find_element(By.XPATH, "/html/body/div/div[2]/div/div/div[3]/div/div/div/div[6]/div[1]/table/tbody/tr/td[3]/p[1]")
        id = first_p.text

        driver.get(f"https://gp.garage.mn/parts/{str(id)}/edit")
        
        link = WebDriverWait(driver, 10).until(
        EC.element_to_be_clickable((By.ID, "select2-genericarticle-container"))
        )
        link.click()
        
        generic_item = driver.find_element(By.XPATH, "/html/body/span/span/span[1]/input")
        generic_item.send_keys(str(Generic_ID))
        time.sleep(1)
        generic_item.send_keys(Keys.ENTER)

        
        # driver.find_element(By.XPATH, '//*[@id="layout-wrapper"]/div[2]/div/div/div[2]/div/div/div/form/div[2]/button').click()
        wait = WebDriverWait(driver, 10)
        button = wait.until(EC.element_to_be_clickable((By.XPATH, '//*[@id="layout-wrapper"]/div[2]/div/div/div[2]/div/div/div/form/div[2]/button')))
        button.click()
    
    else:
        continue
    
    
